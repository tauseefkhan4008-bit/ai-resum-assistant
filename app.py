"""ATS Resume Checker - Streamlit app powered by Google Gemini Flash."""

import io
import json
import os

import streamlit as st
from docx import Document
from google import genai
from google.genai import types

# ---------------------------------------------------------------- config
DEFAULT_MODEL = "gemini-2.5-flash"  # change in the sidebar if Google releases a newer Flash
MAX_FILE_MB = 5

SYSTEM_PROMPT = """You are an expert ATS (Applicant Tracking System) analyst and
professional resume reviewer. Evaluate the resume honestly and strictly.
Never invent information that is not in the resume.

Return ONLY valid JSON with exactly this structure:
{
  "ats_score": <integer 0-100>,
  "summary": "<2-3 sentence overall assessment>",
  "section_scores": {
    "formatting": <0-100>,
    "keywords": <0-100>,
    "experience_impact": <0-100>,
    "skills": <0-100>,
    "education": <0-100>,
    "readability": <0-100>
  },
  "strengths": ["..."],
  "weaknesses": ["..."],
  "improvements": [
    {"priority": "High|Medium|Low", "section": "...", "issue": "...", "fix": "..."}
  ],
  "missing_keywords": ["..."],
  "rewrite_examples": [
    {"original": "...", "improved": "..."}
  ]
}

Scoring guide: ATS-friendly formatting (simple layout, standard headings, no tables
or graphics), relevant keywords, quantified achievements, strong action verbs,
clear contact info, consistent dates, appropriate length.
If a job description is provided, judge keyword match against it and list the
important missing keywords. Otherwise judge for general ATS-readiness.
Give 5-8 improvements ordered by priority and up to 3 rewrite examples that
use only facts already in the resume."""


# ---------------------------------------------------------------- helpers
def get_api_key(sidebar_value: str = "") -> str:
    """Priority: sidebar input > Streamlit secrets > environment variable."""
    if sidebar_value.strip():
        return sidebar_value.strip()
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return st.secrets["GEMINI_API_KEY"]
    except Exception:  # no secrets file locally
        pass
    return os.environ.get("GEMINI_API_KEY", "")


def extract_docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def build_contents(file_name: str, data: bytes, job_description: str):
    """Return the `contents` list for Gemini based on file type."""
    name = file_name.lower()
    task = "Analyze this resume and return the JSON."
    if job_description.strip():
        task += f"\n\nTarget job description:\n{job_description.strip()}"

    if name.endswith(".pdf"):
        # Gemini reads PDFs natively (works for scanned PDFs too)
        return [types.Part.from_bytes(data=data, mime_type="application/pdf"), task]
    if name.endswith(".docx"):
        text = extract_docx_text(data)
    elif name.endswith(".txt"):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Unsupported file type. Upload a PDF, DOCX or TXT file.")

    if len(text.strip()) < 50:
        raise ValueError("Could not read enough text from this file.")
    return [f"RESUME TEXT:\n{text}", task]


def parse_response(raw: str) -> dict:
    """Parse model JSON safely (handles ```json fences) and clamp scores."""
    raw = (raw or "").strip()
    if raw.startswith("```"):
        raw = raw.strip("`")
        if raw.lower().startswith("json"):
            raw = raw[4:]
    start, end = raw.find("{"), raw.rfind("}")
    if start == -1 or end == -1:
        raise ValueError("The model did not return JSON.")
    data = json.loads(raw[start : end + 1])

    def clamp(v):
        try:
            return max(0, min(100, int(round(float(v)))))
        except (TypeError, ValueError):
            return 0

    data["ats_score"] = clamp(data.get("ats_score"))
    data["section_scores"] = {
        k: clamp(v) for k, v in (data.get("section_scores") or {}).items()
    }
    for key in ("strengths", "weaknesses", "improvements", "missing_keywords", "rewrite_examples"):
        if not isinstance(data.get(key), list):
            data[key] = []
    data["summary"] = str(data.get("summary", ""))
    return data


def analyze_resume(api_key: str, model: str, contents) -> dict:
    client = genai.Client(api_key=api_key)
    response = client.models.generate_content(
        model=model,
        contents=contents,
        config=types.GenerateContentConfig(
            system_instruction=SYSTEM_PROMPT,
            response_mime_type="application/json",
            temperature=0.2,
        ),
    )
    return parse_response(response.text)


def score_label(score: int) -> str:
    if score >= 80:
        return "🟢 Excellent"
    if score >= 60:
        return "🟡 Good, needs polish"
    if score >= 40:
        return "🟠 Needs work"
    return "🔴 Poor"


# ---------------------------------------------------------------- UI
def render_results(result: dict):
    score = result["ats_score"]
    st.subheader("Your ATS Score")
    c1, c2 = st.columns([1, 2])
    c1.metric("ATS Score", f"{score}/100")
    c2.write(score_label(score))
    c2.progress(score / 100)
    if result["summary"]:
        st.info(result["summary"])

    if result["section_scores"]:
        st.subheader("Section scores")
        cols = st.columns(3)
        for i, (name, val) in enumerate(result["section_scores"].items()):
            cols[i % 3].metric(name.replace("_", " ").title(), f"{val}/100")

    left, right = st.columns(2)
    with left:
        st.subheader("✅ Strengths")
        for s in result["strengths"]:
            st.markdown(f"- {s}")
    with right:
        st.subheader("⚠️ Weaknesses")
        for w in result["weaknesses"]:
            st.markdown(f"- {w}")

    st.subheader("🛠️ Suggested improvements")
    icons = {"high": "🔴", "medium": "🟠", "low": "🟢"}
    for imp in result["improvements"]:
        if not isinstance(imp, dict):
            continue
        icon = icons.get(str(imp.get("priority", "")).lower(), "⚪")
        title = f"{icon} {imp.get('priority', '')} - {imp.get('section', '')}"
        with st.expander(title):
            st.markdown(f"**Issue:** {imp.get('issue', '')}")
            st.markdown(f"**Fix:** {imp.get('fix', '')}")

    if result["missing_keywords"]:
        st.subheader("🔑 Missing keywords")
        st.write(", ".join(f"`{k}`" for k in result["missing_keywords"]))

    if result["rewrite_examples"]:
        st.subheader("✍️ Rewrite examples")
        for ex in result["rewrite_examples"]:
            if isinstance(ex, dict):
                st.markdown(f"**Before:** {ex.get('original', '')}")
                st.markdown(f"**After:** {ex.get('improved', '')}")
                st.divider()

    st.download_button(
        "Download report (JSON)",
        data=json.dumps(result, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


def main():
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="wide")
    st.title("📄 ATS Resume Checker")
    st.caption("Upload your resume to get an ATS score and concrete improvements.")

    with st.sidebar:
        st.header("Settings")
        key_input = st.text_input("Gemini API key", type="password",
                                  help="Optional if set in Streamlit secrets.")
        model = st.text_input("Model", value=DEFAULT_MODEL)
        st.markdown("[Get a free API key](https://aistudio.google.com/app/apikey)")

    uploaded = st.file_uploader("Upload resume", type=["pdf", "docx", "txt"])
    job_desc = st.text_area("Job description (optional, improves keyword matching)",
                            height=150)

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        api_key = get_api_key(key_input)
        if not api_key:
            st.error("Please provide a Gemini API key in the sidebar or in secrets.")
            return
        data = uploaded.getvalue()
        if len(data) > MAX_FILE_MB * 1024 * 1024:
            st.error(f"File too large. Max {MAX_FILE_MB} MB.")
            return
        try:
            with st.spinner("Analyzing your resume..."):
                contents = build_contents(uploaded.name, data, job_desc)
                st.session_state["result"] = analyze_resume(api_key, model.strip(), contents)
        except ValueError as e:
            st.error(str(e))
            return
        except json.JSONDecodeError:
            st.error("The AI returned an unreadable response. Please try again.")
            return
        except Exception as e:  # API/network errors
            st.error(f"Something went wrong: {e}")
            return

    if "result" in st.session_state:
        render_results(st.session_state["result"])


if __name__ == "__main__":
    main()
